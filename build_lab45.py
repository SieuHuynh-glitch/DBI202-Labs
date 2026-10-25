#!/usr/bin/env python3
"""
Complete Professional Builder for DBI202 Lab 4 and Lab 5.
Student: Huỳnh Hoàng Siêu | Student ID: QE210005 | Class: SE20D
Course: DBI202 - Database Systems | FPT University
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Color palette: Elegant Deep Navy & Slate Theme
COLOR_NAVY = RGBColor(0, 51, 102)        # #003366
COLOR_SLATE = RGBColor(74, 111, 165)     # #4A6FA5
COLOR_DARK = RGBColor(40, 40, 40)        # #282828
COLOR_MUTED = RGBColor(100, 100, 100)    # #646464
COLOR_RED = RGBColor(180, 40, 40)        # #B42828
COLOR_GREEN = RGBColor(34, 139, 34)      # #228B22

HEX_NAVY = "003366"
HEX_SLATE = "4A6FA5"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_LIGHT_GRAY = "F8F9FA"
HEX_BORDER = "D0D7DE"
HEX_HEADER_BG = "003366"

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_block(doc, lab_title, lab_subtitle):
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_after = Pt(2)
    run_uni = p_uni.add_run("FPT UNIVERSITY - DEPARTMENT OF INFORMATION TECHNOLOGY")
    run_uni.font.name = "Calibri"
    run_uni.font.size = Pt(11)
    run_uni.font.bold = True
    run_uni.font.color.rgb = COLOR_SLATE

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.space_after = Pt(6)
    run_course = p_course.add_run("COURSE: DBI202 - DATABASE SYSTEMS")
    run_course.font.name = "Calibri"
    run_course.font.size = Pt(13)
    run_course.font.bold = True
    run_course.font.color.rgb = COLOR_NAVY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(3)
    run_title = p_title.add_run(lab_title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    if lab_subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(10)
        run_sub = p_sub.add_run(lab_subtitle)
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_MUTED

    # Student Info Table
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.5), Inches(2.1), Inches(2.1)]
    info_data = [
        [("Student Name:", True), ("Student ID:", True), ("Class:", True)],
        [("Huỳnh Hoàng Siêu", False), ("QE210005", False), ("SE20D", False)]
    ]

    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            text, is_label = info_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if is_label:
                run.font.bold = True
                run.font.color.rgb = COLOR_SLATE
                set_cell_shading(cell, HEX_LIGHT_BLUE)
            else:
                run.font.bold = True
                run.font.color.rgb = COLOR_NAVY
                set_cell_shading(cell, HEX_LIGHT_GRAY)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    set_table_borders(table, color="B0C4DE", sz="6", val="single")

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(12)
    p_div.paragraph_format.space_after = Pt(12)
    p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{HEX_NAVY}"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE
    return h

def add_callout(doc, title, text, bg_hex=HEX_LIGHT_BLUE, border_color="003366"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.7)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=70, bottom=70, left=130, right=130)
    
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if title:
        run_t = p.add_run(title + "\n")
        run_t.font.name = "Calibri"
        run_t.font.bold = True
        run_t.font.size = Pt(9.5)
        run_t.font.color.rgb = COLOR_NAVY
    
    run_b = p.add_run(text)
    run_b.font.name = "Consolas"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = COLOR_DARK

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def style_standard_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=50, bottom=50, left=80, right=80)

    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = HEX_LIGHT_GRAY if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].width = col_widths[c_idx]
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or len(str(val)) < 8 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_DARK
            set_cell_shading(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=40, bottom=40, left=80, right=80)

    set_table_borders(table, color="D0D7DE", sz="4", val="single")

# ==========================================
# LAB 4: 52 SQL QUERIES (COMPANYDB)
# ==========================================
def generate_lab4(output_path="output/Lab4.docx"):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    add_header_block(doc, "LAB 4: ADVANCED SQL QUERIES", "Complex Relational Joins, Aggregations, Grouping & Subqueries on CompanyDB")

    add_heading_1(doc, "1. Database Schema Overview (CompanyDB)")
    p_desc = doc.add_paragraph()
    p_desc.add_run("This lab requires implementing 52 comprehensive SQL queries on the standard enterprise CompanyDB database. The relational physical schema comprises seven core tables:")

    add_callout(doc, "CompanyDB Relational Schema Definition:",
                "1. tblDepartment (depNum [PK], depName, mgrSSN [FK → tblEmployee], mgrStartdate)\n"
                "2. tblEmployee (empSSN [PK], empName, empBirthdate, empSex, empSalary, empAddress,\n"
                "                depNum [FK → tblDepartment], supervisorSSN [FK → tblEmployee], empStartdate)\n"
                "3. tblProject (proNum [PK], proName, locNum [FK → tblLocation], depNum [FK → tblDepartment])\n"
                "4. tblLocation (locNum [PK], locName)\n"
                "5. tblDepLocation (depNum [PK, FK → tblDepartment], locNum [PK, FK → tblLocation])\n"
                "6. tblWorksOn (empSSN [PK, FK → tblEmployee], proNum [PK, FK → tblProject], workHours)\n"
                "7. tblDependent (depName [PK], empSSN [PK, FK → tblEmployee], depSex, depBirthdate, depRelationship)")

    add_heading_1(doc, "2. Solutions for 52 SQL Queries")

    raw_queries = [
        # 1-10
        (1, "Cho biết ai đang quản lý phòng ban có tên: Phòng Nghiên cứu và phát triển. Thông tin yêu cầu: mã số, họ tên nhân viên, mã số phòng ban, tên phòng ban",
         "Find the manager of department 'Phòng Nghiên cứu và phát triển'. Required: empSSN, empName, depNum, depName",
         "Join tblDepartment and tblEmployee on d.mgrSSN = e.empSSN where depName matches.",
         "SELECT e.empSSN, e.empName, d.depNum, d.depName\nFROM tblDepartment d\nINNER JOIN tblEmployee e ON d.mgrSSN = e.empSSN\nWHERE d.depName = N'Phòng Nghiên cứu và phát triển';"),

        (2, "Cho phòng ban có tên: Phòng Nghiên cứu và phát triển hiện đang quản lý dự án nào. Thông tin yêu cầu: mã số dụ án, tên dự án, tên phòng ban quản lý",
         "Find projects managed by 'Phòng Nghiên cứu và phát triển'. Required: proNum, proName, depName",
         "Join tblProject and tblDepartment on p.depNum = d.depNum filtered by department name.",
         "SELECT p.proNum, p.proName, d.depName\nFROM tblProject p\nINNER JOIN tblDepartment d ON p.depNum = d.depNum\nWHERE d.depName = N'Phòng Nghiên cứu và phát triển';"),

        (3, "Cho biết dự án có tên ProjectB hiện đang được quản lý bởi phòng ban nào. Thông tin yêu cầu: mã số dụ án, tên dự án, tên phòng ban quản lý",
         "Find the managing department of project 'ProjectB'. Required: proNum, proName, depName",
         "Join tblProject with tblDepartment where proName equals 'ProjectB'.",
         "SELECT p.proNum, p.proName, d.depName\nFROM tblProject p\nINNER JOIN tblDepartment d ON p.depNum = d.depNum\nWHERE p.proName = 'ProjectB';"),

        (4, "Cho biết những nhân viên nào đang bị giám sát bởi nhân viên có tên Mai Duy An. Thông tin yêu cầu: mã số nhân viên, họ tên nhân viên",
         "Find employees supervised by 'Mai Duy An'. Required: empSSN, empName",
         "Self-join tblEmployee on e.supervisorSSN = s.empSSN where s.empName = N'Mai Duy An'.",
         "SELECT e.empSSN, e.empName\nFROM tblEmployee e\nINNER JOIN tblEmployee s ON e.supervisorSSN = s.empSSN\nWHERE s.empName = N'Mai Duy An';"),

        (5, "Cho biết ai hiện đang giám sát những nhân viên có tên Mai Duy An. Thông tin yêu cầu: mã số nhân viên, họ tên nhân viên giám sát.",
         "Find the supervisor of employee 'Mai Duy An'. Required: empSSN, supervisorName",
         "Self-join tblEmployee where the supervised employee's name is 'Mai Duy An'.",
         "SELECT s.empSSN, s.empName AS supervisorName\nFROM tblEmployee e\nINNER JOIN tblEmployee s ON e.supervisorSSN = s.empSSN\nWHERE e.empName = N'Mai Duy An';"),

        (6, "Cho biết dự án có tên ProjectA hiện đang làm việc ở đâu. Thông tin yêu cầu: mã số, tên vị trí làm việc.",
         "Find location of project 'ProjectA'. Required: locNum, locName",
         "Join tblProject and tblLocation on p.locNum = l.locNum where proName = 'ProjectA'.",
         "SELECT l.locNum, l.locName\nFROM tblProject p\nINNER JOIN tblLocation l ON p.locNum = l.locNum\nWHERE p.proName = 'ProjectA';"),

        (7, "Cho biết vị trí làm việc có tên Tp. HCM hiện đang là chỗ làm việc của những dự án nào. Thông tin yêu cầu: mã số, tên dự án",
         "Find projects located in 'Tp. HCM'. Required: proNum, proName",
         "Join tblProject and tblLocation where locName = N'Tp. HCM'.",
         "SELECT p.proNum, p.proName\nFROM tblProject p\nINNER JOIN tblLocation l ON p.locNum = l.locNum\nWHERE l.locName = N'Tp. HCM';"),

        (8, "Cho biết những người phụ thuộc trên 18 tuổi. Thông tin yêu cầu: tên, ngày tháng năm sinh của người phụ thuộc, tên nhân viên phụ thuộc vào.",
         "Find dependents older than 18 years old. Required: depName, depBirthdate, empName",
         "Join tblDependent and tblEmployee; a dependent is older than 18 when the 18th birthday has already passed: DATEADD(YEAR, 18, depBirthdate) < GETDATE(). (DATEDIFF(YEAR, ...) only subtracts calendar years and ignores month/day, so it gives wrong ages.)",
         "SELECT d.depName, d.depBirthdate, e.empName\nFROM tblDependent d\nINNER JOIN tblEmployee e ON d.empSSN = e.empSSN\nWHERE DATEADD(YEAR, 18, d.depBirthdate) < GETDATE();"),

        (9, "Cho biết những người phụ thuộc là nam giới. Thông tin yêu cầu: tên, ngày tháng năm sinh của người phụ thuộc, tên nhân viên phụ thuộc vào",
         "Find male dependents. Required: depName, depBirthdate, empName",
         "Join tblDependent and tblEmployee filtering depSex for male ('M' or N'Nam').",
         "SELECT d.depName, d.depBirthdate, e.empName\nFROM tblDependent d\nINNER JOIN tblEmployee e ON d.empSSN = e.empSSN\nWHERE d.depSex = 'M' OR d.depSex = N'Nam';"),

        (10, "Cho biết những nơi làm việc của phòng ban có tên: Phòng Nghiên cứu và phát triển. Thông tin yêu cầu: mã phòng ban, tên phòng ban, tên nơi làm việc.",
         "Find locations of department 'Phòng Nghiên cứu và phát triển'. Required: depNum, depName, locName",
         "Join tblDepartment, tblDepLocation, and tblLocation.",
         "SELECT d.depNum, d.depName, l.locName\nFROM tblDepartment d\nINNER JOIN tblDepLocation dl ON d.depNum = dl.depNum\nINNER JOIN tblLocation l ON dl.locNum = l.locNum\nWHERE d.depName = N'Phòng Nghiên cứu và phát triển';"),

        # 11-20
        (11, "Cho biết các dự án làm việc tại Tp. HCM. Thông tin yêu cầu: mã dự án, tên dự án, tên phòng ban chịu trách nhiệm dự án.",
         "Find projects located in 'Tp. HCM' and managing department. Required: proNum, proName, depName",
         "Join tblProject, tblLocation, and tblDepartment.",
         "SELECT p.proNum, p.proName, d.depName\nFROM tblProject p\nINNER JOIN tblLocation l ON p.locNum = l.locNum\nINNER JOIN tblDepartment d ON p.depNum = d.depNum\nWHERE l.locName = N'Tp. HCM';"),

        (12, "Cho biết những người phụ thuộc là nữ giới, của nhân viên thuộc phòng ban có tên: Phòng Nghiên cứu và phát triển. Thông tin yêu cầu: tên nhân viên, tên người phụ thuộc, mối liên hệ",
         "Find female dependents of employees in 'Phòng Nghiên cứu và phát triển'. Required: empName, depName, depRelationship",
         "Join tblEmployee, tblDepartment, and tblDependent filtering female depSex.",
         "SELECT e.empName, d.depName, d.depRelationship\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblDependent d ON e.empSSN = d.empSSN\nWHERE dep.depName = N'Phòng Nghiên cứu và phát triển'\n  AND (d.depSex = 'F' OR d.depSex = N'Nữ');"),

        (13, "Cho biết những người phụ thuộc trên 18 tuổi, của nhân viên thuộc phòng ban có tên: Phòng Nghiên cứu và phát triển. Thông tin yêu cầu: tên nhân viên, tên người phụ thuộc, mối liên hệ",
         "Find dependents over 18 of employees in 'Phòng Nghiên cứu và phát triển'. Required: empName, depName, depRelationship",
         "Filter by department name and DATEADD(YEAR, 18, depBirthdate) < GETDATE() (18th birthday already passed; DATEDIFF(YEAR) ignores month/day so it is inaccurate).",
         "SELECT e.empName, d.depName, d.depRelationship\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblDependent d ON e.empSSN = d.empSSN\nWHERE dep.depName = N'Phòng Nghiên cứu và phát triển'\n  AND DATEADD(YEAR, 18, d.depBirthdate) < GETDATE();"),

        (14, "Cho biết số lượng người phụ thuộc theo giới tính. Thông tin yêu cầu: giới tính, số lượng người phụ thuộc",
         "Count number of dependents by gender. Required: depSex, count",
         "GROUP BY depSex and COUNT(*).",
         "SELECT depSex, COUNT(*) AS numberOfDependents\nFROM tblDependent\nGROUP BY depSex;"),

        (15, "Cho biết số lượng người phụ thuộc theo mối liên hệ với nhân viên. Thông tin yêu cầu: mối liên hệ, số lượng người phụ thuộc",
         "Count number of dependents by relationship. Required: depRelationship, count",
         "GROUP BY depRelationship and COUNT(*).",
         "SELECT depRelationship, COUNT(*) AS numberOfDependents\nFROM tblDependent\nGROUP BY depRelationship;"),

        (16, "Cho biết số lượng người phụ thuộc theo mỗi phòng ban. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng người phụ thuộc",
         "Count number of dependents per department. Required: depNum, depName, count",
         "LEFT JOIN tblDepartment, tblEmployee, tblDependent so every department is listed (0 when it has no dependents), COUNT(d.depName) GROUP BY depNum, depName.",
         "SELECT dep.depNum, dep.depName, COUNT(d.depName) AS numberOfDependents\nFROM tblDepartment dep\nLEFT JOIN tblEmployee e ON dep.depNum = e.depNum\nLEFT JOIN tblDependent d ON e.empSSN = d.empSSN\nGROUP BY dep.depNum, dep.depName;"),

        (17, "Cho biết phòng ban nào có số lượng người phụ thuộc là ít nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng người phụ thuộc",
         "Find department with minimum dependents. Required: depNum, depName, count",
         "Use LEFT JOIN to include departments with 0 dependents, TOP 1 WITH TIES ORDER BY count ASC.",
         "SELECT TOP 1 WITH TIES dep.depNum, dep.depName, COUNT(d.depName) AS numberOfDependents\nFROM tblDepartment dep\nLEFT JOIN tblEmployee e ON dep.depNum = e.depNum\nLEFT JOIN tblDependent d ON e.empSSN = d.empSSN\nGROUP BY dep.depNum, dep.depName\nORDER BY numberOfDependents ASC;"),

        (18, "Cho biết phòng ban nào có số lượng người phụ thuộc là nhiều nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng người phụ thuộc",
         "Find department with maximum dependents. Required: depNum, depName, count",
         "TOP 1 WITH TIES ORDER BY count DESC.",
         "SELECT TOP 1 WITH TIES dep.depNum, dep.depName, COUNT(d.depName) AS numberOfDependents\nFROM tblDepartment dep\nINNER JOIN tblEmployee e ON dep.depNum = e.depNum\nINNER JOIN tblDependent d ON e.empSSN = d.empSSN\nGROUP BY dep.depNum, dep.depName\nORDER BY numberOfDependents DESC;"),

        (19, "Cho biết tổng số giờ tham gia dự án của mỗi nhân viên. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tên phòng ban của nhân viên",
         "Find total project work hours per employee. Required: empSSN, empName, depName, totalHours",
         "Join tblEmployee, tblDepartment, tblWorksOn, SUM(workHours), GROUP BY.",
         "SELECT e.empSSN, e.empName, dep.depName, SUM(w.workHours) AS totalWorkHours\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName, dep.depName;"),

        (20, "Cho biết tổng số giờ làm dự án của mỗi phòng ban. Thông tin yêu cầu: mã phòng ban, tên phòng ban, tổng số giờ",
         "Find total project work hours per department. Required: depNum, depName, totalHours",
         "Join tblDepartment, tblEmployee, tblWorksOn, SUM(workHours) GROUP BY department.",
         "SELECT dep.depNum, dep.depName, SUM(w.workHours) AS totalProjectHours\nFROM tblDepartment dep\nINNER JOIN tblEmployee e ON dep.depNum = e.depNum\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY dep.depNum, dep.depName;"),

        # 21-30
        (21, "Cho biết nhân viên nào có số giờ tham gia dự án là ít nhất. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tổng số giờ tham gia dự án",
         "Find employee with least project work hours. Required: empSSN, empName, totalHours",
         "TOP 1 WITH TIES ordered by SUM(workHours) ASC.",
         "SELECT TOP 1 WITH TIES e.empSSN, e.empName, SUM(w.workHours) AS totalHours\nFROM tblEmployee e\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName\nORDER BY totalHours ASC;"),

        (22, "Cho biết nhân viên nào có số giờ tham gia dự án là nhiều nhất. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tổng số giờ tham gia dự án",
         "Find employee with most project work hours. Required: empSSN, empName, totalHours",
         "TOP 1 WITH TIES ordered by SUM(workHours) DESC.",
         "SELECT TOP 1 WITH TIES e.empSSN, e.empName, SUM(w.workHours) AS totalHours\nFROM tblEmployee e\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName\nORDER BY totalHours DESC;"),

        (23, "Cho biết những nhân viên nào lần đầu tiên tham gia dụ án. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tên phòng ban của nhân viên",
         "Find employees participating in exactly 1 project (first time). Required: empSSN, empName, depName",
         "HAVING COUNT(w.proNum) = 1.",
         "SELECT e.empSSN, e.empName, dep.depName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName, dep.depName\nHAVING COUNT(w.proNum) = 1;"),

        (24, "Cho biết những nhân viên nào lần thứ hai tham gia dụ án. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tên phòng ban của nhân viên",
         "Find employees participating in exactly 2 projects (second time). Required: empSSN, empName, depName",
         "HAVING COUNT(w.proNum) = 2.",
         "SELECT e.empSSN, e.empName, dep.depName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName, dep.depName\nHAVING COUNT(w.proNum) = 2;"),

        (25, "Cho biết những nhân viên nào tham gia tối thiểu hai dụ án. Thông tin yêu cầu: mã nhân viên, tên nhân viên, tên phòng ban của nhân viên",
         "Find employees participating in at least 2 projects. Required: empSSN, empName, depName",
         "HAVING COUNT(w.proNum) >= 2.",
         "SELECT e.empSSN, e.empName, dep.depName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nGROUP BY e.empSSN, e.empName, dep.depName\nHAVING COUNT(w.proNum) >= 2;"),

        (26, "Cho biết số lượng thành viên của mỗi dự án. Thông tin yêu cầu: mã dự án, tên dự án, số lượng thành viên",
         "Find number of members per project. Required: proNum, proName, memberCount",
         "Join tblProject and tblWorksOn, COUNT(empSSN), GROUP BY proNum, proName.",
         "SELECT p.proNum, p.proName, COUNT(w.empSSN) AS numberOfMembers\nFROM tblProject p\nLEFT JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName;"),

        (27, "Cho biết tổng số giờ làm của mỗi dự án. Thông tin yêu cầu: mã dự án, tên dự án, tổng số giờ làm",
         "Find total work hours of each project. Required: proNum, proName, totalHours",
         "Join tblProject and tblWorksOn, SUM(workHours), GROUP BY proNum, proName.",
         "SELECT p.proNum, p.proName, SUM(w.workHours) AS totalProjectHours\nFROM tblProject p\nINNER JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName;"),

        (28, "Cho biết dự án nào có số lượng thành viên là ít nhất. Thông tin yêu cầu: mã dự án, tên dự án, số lượng thành viên",
         "Find project with fewest members. Required: proNum, proName, memberCount",
         "LEFT JOIN so projects with 0 members are included, TOP 1 WITH TIES ordered by COUNT(empSSN) ASC.",
         "SELECT TOP 1 WITH TIES p.proNum, p.proName, COUNT(w.empSSN) AS numberOfMembers\nFROM tblProject p\nLEFT JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName\nORDER BY numberOfMembers ASC;"),

        (29, "Cho biết dự án nào có số lượng thành viên là nhiều nhất. Thông tin yêu cầu: mã dự án, tên dự án, số lượng thành viên",
         "Find project with most members. Required: proNum, proName, memberCount",
         "TOP 1 WITH TIES ordered by COUNT(empSSN) DESC.",
         "SELECT TOP 1 WITH TIES p.proNum, p.proName, COUNT(w.empSSN) AS numberOfMembers\nFROM tblProject p\nINNER JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName\nORDER BY numberOfMembers DESC;"),

        (30, "Cho biết dự án nào có tổng số giờ làm là ít nhất. Thông tin yêu cầu: mã dự án, tên dự án, tổng số giờ làm",
         "Find project with least total work hours. Required: proNum, proName, totalHours",
         "TOP 1 WITH TIES ordered by SUM(workHours) ASC.",
         "SELECT TOP 1 WITH TIES p.proNum, p.proName, SUM(w.workHours) AS totalHours\nFROM tblProject p\nINNER JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName\nORDER BY totalHours ASC;"),

        # 31-40
        (31, "Cho biết dự án nào có tổng số giờ làm là nhiều nhất. Thông tin yêu cầu: mã dự án, tên dự án, tổng số giờ làm",
         "Find project with most total work hours. Required: proNum, proName, totalHours",
         "TOP 1 WITH TIES ordered by SUM(workHours) DESC.",
         "SELECT TOP 1 WITH TIES p.proNum, p.proName, SUM(w.workHours) AS totalHours\nFROM tblProject p\nINNER JOIN tblWorksOn w ON p.proNum = w.proNum\nGROUP BY p.proNum, p.proName\nORDER BY totalHours DESC;"),

        (32, "Cho biết số lượng phòng ban làm việc theo mỗi nơi làm việc. Thông tin yêu cầu: tên nơi làm việc, số lượng phòng ban",
         "Find count of departments working per location. Required: locName, departmentCount",
         "LEFT JOIN tblLocation and tblDepLocation so every location is listed (0 when no department works there), COUNT(dl.depNum) GROUP BY locName.",
         "SELECT l.locName, COUNT(dl.depNum) AS numberOfDepartments\nFROM tblLocation l\nLEFT JOIN tblDepLocation dl ON l.locNum = dl.locNum\nGROUP BY l.locName;"),

        (33, "Cho biết số lượng chỗ làm việc theo mỗi phòng ban. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng chỗ làm việc",
         "Find number of locations per department. Required: depNum, depName, locationCount",
         "LEFT JOIN tblDepartment and tblDepLocation so every department is listed (0 when it has no location), COUNT(dl.locNum) GROUP BY depNum, depName.",
         "SELECT d.depNum, d.depName, COUNT(dl.locNum) AS numberOfLocations\nFROM tblDepartment d\nLEFT JOIN tblDepLocation dl ON d.depNum = dl.depNum\nGROUP BY d.depNum, d.depName;"),

        (34, "Cho biết phòng ban nào có nhiều chỗ làm việc nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng chỗ làm việc",
         "Find department with most locations. Required: depNum, depName, locationCount",
         "TOP 1 WITH TIES ordered by COUNT(locNum) DESC.",
         "SELECT TOP 1 WITH TIES d.depNum, d.depName, COUNT(dl.locNum) AS numberOfLocations\nFROM tblDepartment d\nINNER JOIN tblDepLocation dl ON d.depNum = dl.depNum\nGROUP BY d.depNum, d.depName\nORDER BY numberOfLocations DESC;"),

        (35, "Cho biết phòng ban nào có it chỗ làm việc nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng chỗ làm việc",
         "Find department with fewest locations. Required: depNum, depName, locationCount",
         "LEFT JOIN so departments with 0 locations are included, TOP 1 WITH TIES ordered by COUNT(locNum) ASC.",
         "SELECT TOP 1 WITH TIES d.depNum, d.depName, COUNT(dl.locNum) AS numberOfLocations\nFROM tblDepartment d\nLEFT JOIN tblDepLocation dl ON d.depNum = dl.depNum\nGROUP BY d.depNum, d.depName\nORDER BY numberOfLocations ASC;"),

        (36, "Cho biết địa điểm nào có nhiều phòng ban làm việc nhất. Thông tin yêu cầu: tên nơi làm việc, số lượng phòng ban",
         "Find location with most departments working. Required: locName, departmentCount",
         "TOP 1 WITH TIES ordered by COUNT(depNum) DESC.",
         "SELECT TOP 1 WITH TIES l.locName, COUNT(dl.depNum) AS numberOfDepartments\nFROM tblLocation l\nINNER JOIN tblDepLocation dl ON l.locNum = dl.locNum\nGROUP BY l.locName\nORDER BY numberOfDepartments DESC;"),

        (37, "Cho biết địa điểm nào có ít phòng ban làm việc nhất. Thông tin yêu cầu: tên nơi làm việc, số lượng phòng ban",
         "Find location with fewest departments working. Required: locName, departmentCount",
         "LEFT JOIN so locations with 0 departments are included, TOP 1 WITH TIES ordered by COUNT(depNum) ASC.",
         "SELECT TOP 1 WITH TIES l.locName, COUNT(dl.depNum) AS numberOfDepartments\nFROM tblLocation l\nLEFT JOIN tblDepLocation dl ON l.locNum = dl.locNum\nGROUP BY l.locName\nORDER BY numberOfDepartments ASC;"),

        (38, "Cho biết nhân viên nào có nhiều người phụ thuộc nhất. Thông tin yêu cầu: mã số, họ tên nhân viên, số lượng người phụ thuộc",
         "Find employee with most dependents. Required: empSSN, empName, dependentCount",
         "Join tblEmployee and tblDependent, TOP 1 WITH TIES ORDER BY COUNT(depName) DESC.",
         "SELECT TOP 1 WITH TIES e.empSSN, e.empName, COUNT(d.depName) AS numberOfDependents\nFROM tblEmployee e\nINNER JOIN tblDependent d ON e.empSSN = d.empSSN\nGROUP BY e.empSSN, e.empName\nORDER BY numberOfDependents DESC;"),

        (39, "Cho biết nhân viên nào có ít người phụ thuộc nhất. Thông tin yêu cầu: mã số, họ tên nhân viên, số lượng người phụ thuộc",
         "Find employee with least dependents (among those with dependents). Required: empSSN, empName, dependentCount",
         "TOP 1 WITH TIES ORDER BY COUNT(depName) ASC.",
         "SELECT TOP 1 WITH TIES e.empSSN, e.empName, COUNT(d.depName) AS numberOfDependents\nFROM tblEmployee e\nINNER JOIN tblDependent d ON e.empSSN = d.empSSN\nGROUP BY e.empSSN, e.empName\nORDER BY numberOfDependents ASC;"),

        (40, "Cho biết nhân viên nào không có người phụ thuộc. Thông tin yêu cầu: mã số nhân viên, họ tên nhân viên, tên phòng ban của nhân viên",
         "Find employees who have NO dependents. Required: empSSN, empName, depName",
         "Use NOT IN or NOT EXISTS on tblDependent.",
         "SELECT e.empSSN, e.empName, dep.depName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nWHERE e.empSSN NOT IN (SELECT empSSN FROM tblDependent);"),

        # 41-52
        (41, "Cho biết phòng ban nào không có người phụ thuộc. Thông tin yêu cầu: mã số phòng ban, tên phòng ban",
         "Find departments where no employee has dependents. Required: depNum, depName",
         "NOT EXISTS: no employee of the department has a dependent (NULL-safe, unlike NOT IN when e.depNum can be NULL).",
         "SELECT dep.depNum, dep.depName\nFROM tblDepartment dep\nWHERE NOT EXISTS (\n    SELECT 1\n    FROM tblEmployee e\n    INNER JOIN tblDependent d ON e.empSSN = d.empSSN\n    WHERE e.depNum = dep.depNum\n);"),

        (42, "Cho biết những nhân viên nào chưa hề tham gia vào bất kỳ dự án nào. Thông tin yêu cầu: mã số, tên nhân viên, tên phòng ban của nhân viên",
         "Find employees who have never worked on any project. Required: empSSN, empName, depName",
         "NOT IN (SELECT empSSN FROM tblWorksOn).",
         "SELECT e.empSSN, e.empName, dep.depName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nWHERE e.empSSN NOT IN (SELECT DISTINCT empSSN FROM tblWorksOn);"),

        (43, "Cho biết phòng ban không có nhân viên nào tham gia (bất kỳ) dự án. Thông tin yêu cầu: mã số phòng ban, tên phòng ban",
         "Find departments with no employees participating in any project. Required: depNum, depName",
         "NOT EXISTS: no employee of the department appears in tblWorksOn (NULL-safe).",
         "SELECT dep.depNum, dep.depName\nFROM tblDepartment dep\nWHERE NOT EXISTS (\n    SELECT 1\n    FROM tblEmployee e\n    INNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\n    WHERE e.depNum = dep.depNum\n);"),

        (44, "Cho biết phòng ban không có nhân viên nào tham gia vào dự án có tên là ProjectA. Thông tin yêu cầu: mã số phòng ban, tên phòng ban",
         "Find departments where no employee works on 'ProjectA'. Required: depNum, depName",
         "NOT EXISTS over employee, workson, project filtering proName = 'ProjectA' (NULL-safe).",
         "SELECT dep.depNum, dep.depName\nFROM tblDepartment dep\nWHERE NOT EXISTS (\n    SELECT 1\n    FROM tblEmployee e\n    INNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\n    INNER JOIN tblProject p ON w.proNum = p.proNum\n    WHERE e.depNum = dep.depNum\n      AND p.proName = 'ProjectA'\n);"),

        (45, "Cho biết số lượng dự án được quản lý theo mỗi phòng ban. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng dự án",
         "Find number of projects managed per department. Required: depNum, depName, projectCount",
         "LEFT JOIN tblProject on depNum, COUNT(proNum) GROUP BY depNum, depName.",
         "SELECT dep.depNum, dep.depName, COUNT(p.proNum) AS numberOfProjects\nFROM tblDepartment dep\nLEFT JOIN tblProject p ON dep.depNum = p.depNum\nGROUP BY dep.depNum, dep.depName;"),

        (46, "Cho biết phòng ban nào quản lý it dự án nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng dự án",
         "Find department managing fewest projects. Required: depNum, depName, projectCount",
         "LEFT JOIN, TOP 1 WITH TIES ORDER BY COUNT(proNum) ASC.",
         "SELECT TOP 1 WITH TIES dep.depNum, dep.depName, COUNT(p.proNum) AS numberOfProjects\nFROM tblDepartment dep\nLEFT JOIN tblProject p ON dep.depNum = p.depNum\nGROUP BY dep.depNum, dep.depName\nORDER BY numberOfProjects ASC;"),

        (47, "Cho biết phòng ban nào quản lý nhiều dự án nhất. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng dự án",
         "Find department managing most projects. Required: depNum, depName, projectCount",
         "TOP 1 WITH TIES ORDER BY COUNT(proNum) DESC.",
         "SELECT TOP 1 WITH TIES dep.depNum, dep.depName, COUNT(p.proNum) AS numberOfProjects\nFROM tblDepartment dep\nINNER JOIN tblProject p ON dep.depNum = p.depNum\nGROUP BY dep.depNum, dep.depName\nORDER BY numberOfProjects DESC;"),

        (48, "Cho biết những phòng ban nào có nhiểu hơn 5 nhân viên đang quản lý dự án gì. Thông tin yêu cầu: mã phòng ban, tên phòng ban, số lượng nhân viên của phòng ban, tên dự án quản lý",
         "Find projects managed by departments that have more than 5 employees. Required: depNum, depName, employeeCount, proName",
         "Join department, employee (count > 5 via subquery/CTE or group by), and project.",
         "SELECT dep.depNum, dep.depName, COUNT(DISTINCT e.empSSN) AS numberOfEmployees, p.proName\nFROM tblDepartment dep\nINNER JOIN tblEmployee e ON dep.depNum = e.depNum\nINNER JOIN tblProject p ON dep.depNum = p.depNum\nGROUP BY dep.depNum, dep.depName, p.proName\nHAVING COUNT(DISTINCT e.empSSN) > 5;"),

        (49, "Cho biết những nhân viên thuộc phòng có tên là Phòng nghiên cứu, và không có người phụ thuộc. Thông tin yêu cầu: mã nhân viên, họ tên nhân viên",
         "Find employees in 'Phòng Nghiên cứu' with NO dependents. Required: empSSN, empName",
         "Department name filter with LIKE N'%Phòng Nghiên cứu%' AND NOT IN tblDependent.",
         "SELECT e.empSSN, e.empName\nFROM tblEmployee e\nINNER JOIN tblDepartment dep ON e.depNum = dep.depNum\nWHERE dep.depName LIKE N'%Phòng Nghiên cứu%'\n  AND e.empSSN NOT IN (SELECT empSSN FROM tblDependent);"),

        (50, "Cho biết tổng số giờ làm của các nhân viên, mà các nhân viên này không có người phụ thuộc. Thông tin yêu cầu: mã nhân viên, họ tên nhân viên, tổng số giờ làm",
         "Find total work hours of employees who have NO dependents. Required: empSSN, empName, totalHours",
         "Join tblEmployee, tblWorksOn where empSSN NOT IN tblDependent.",
         "SELECT e.empSSN, e.empName, SUM(w.workHours) AS totalHours\nFROM tblEmployee e\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nWHERE e.empSSN NOT IN (SELECT empSSN FROM tblDependent)\nGROUP BY e.empSSN, e.empName;"),

        (51, "Cho biết tổng số giờ làm của các nhân viên, mà các nhân viên này có nhiều hơn 3 người phụ thuộc. Thông tin yêu cầu: mã nhân viên, họ tên nhân viên, số lượng người phụ thuộc, tổng số giờ làm",
         "Find total work hours of employees with more than 3 dependents. Required: empSSN, empName, dependentCount, totalHours",
         "Aggregate dependents and work hours in separate derived tables, then join: joining tblDependent and tblWorksOn directly multiplies the rows and inflates SUM(workHours) by the number of dependents.",
         "SELECT e.empSSN, e.empName, dc.numberOfDependents, wh.totalHours\nFROM tblEmployee e\nINNER JOIN (SELECT empSSN, COUNT(*) AS numberOfDependents\n            FROM tblDependent\n            GROUP BY empSSN\n            HAVING COUNT(*) > 3) dc ON e.empSSN = dc.empSSN\nINNER JOIN (SELECT empSSN, SUM(workHours) AS totalHours\n            FROM tblWorksOn\n            GROUP BY empSSN) wh ON e.empSSN = wh.empSSN;"),

        (52, "Cho biết tổng số giờ làm việc của các nhân viên hiện đang dưới quyền giám sát (bị quản lý bởi) của nhân viên Mai Duy An. Thông tin yêu cầu: mã nhân viên, họ tên nhân viên, tổng số giờ làm",
         "Find total work hours of employees supervised by 'Mai Duy An'. Required: empSSN, empName, totalHours",
         "Self-join on supervisorSSN where supervisor name is 'Mai Duy An', join tblWorksOn, SUM(workHours).",
         "SELECT e.empSSN, e.empName, SUM(w.workHours) AS totalHours\nFROM tblEmployee e\nINNER JOIN tblEmployee s ON e.supervisorSSN = s.empSSN\nINNER JOIN tblWorksOn w ON e.empSSN = w.empSSN\nWHERE s.empName = N'Mai Duy An'\nGROUP BY e.empSSN, e.empName;")
    ]

    for qid, qvn, qen, qreason, qsql in raw_queries:
        add_heading_2(doc, f"Query {qid}: {qen}")
        
        p_orig = doc.add_paragraph()
        r_lbl = p_orig.add_run("Original Task (VN): ")
        r_lbl.bold = True
        r_lbl.font.color.rgb = COLOR_SLATE
        p_orig.add_run(qvn)
        p_orig.paragraph_format.space_after = Pt(2)

        add_callout(doc, f"SQL Implementation for Query {qid}:", qsql, bg_hex=HEX_LIGHT_BLUE, border_color="003366")

        p_exp = doc.add_paragraph()
        r_exp_lbl = p_exp.add_run("Technical Explanation: ")
        r_exp_lbl.bold = True
        r_exp_lbl.font.color.rgb = COLOR_SLATE
        p_exp.add_run(qreason)
        p_exp.paragraph_format.space_after = Pt(8)

    doc.save(output_path)
    print(f"Lab 4 successfully generated at: {output_path}")

# ==========================================
# LAB 5: DATABASE PROGRAMMING (ABCCOMPANY)
# ==========================================
def generate_lab5(output_path="output/Lab5.docx"):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    add_header_block(doc, "LAB 5: DATABASE PROGRAMMING ON SQL SERVER", "User-Defined Functions, Stored Procedures, DML Triggers & Views on ABCCompany")

    add_heading_1(doc, "1. Database Setup Script (ABCCompany)")
    p_setup = doc.add_paragraph()
    p_setup.add_run("The database ABCCompany manages materials, customers, sales invoices, and line-item details. The complete schema creation script and sample datasets provided by the course are executed below:")

    add_callout(doc, "ABCCompany Schema & Sample Data (course-provided script):",
                "CREATE DATABASE ABCCompany;\n"
                "GO\n"
                "USE ABCCompany;\n"
                "GO\n\n"
                "CREATE TABLE KHACHHANG (\n"
                "    MaKH nvarchar(5) NOT NULL PRIMARY KEY,\n"
                "    TenKH nvarchar(30) NOT NULL,\n"
                "    DiaChi nvarchar(50),\n"
                "    DT nvarchar(10),\n"
                "    Email nvarchar(30)\n"
                ");\n\n"
                "CREATE TABLE VATTU (\n"
                "    MaVT nvarchar(5) NOT NULL PRIMARY KEY,\n"
                "    TenVT nvarchar(30) NOT NULL,\n"
                "    DVT nvarchar(20),\n"
                "    GiaMua int CHECK (GiaMua > 0),\n"
                "    SLTon int CHECK (SLTon >= 0)\n"
                ");\n\n"
                "CREATE TABLE HOADON (\n"
                "    MaHD nvarchar(10) NOT NULL PRIMARY KEY,\n"
                "    Ngay datetime CHECK (Ngay < GETDATE()),\n"
                "    MaKH nvarchar(5) FOREIGN KEY REFERENCES KHACHHANG(MaKH),\n"
                "    TongTG int\n"
                ");\n\n"
                "CREATE TABLE CHITIETHOADON (\n"
                "    MaHD nvarchar(10) FOREIGN KEY REFERENCES HOADON(MaHD),\n"
                "    MaVT nvarchar(5) FOREIGN KEY REFERENCES VATTU(MaVT),\n"
                "    PRIMARY KEY (MaHD, MaVT),\n"
                "    SL int,\n"
                "    KhuyenMai int,\n"
                "    GiaBan int\n"
                ");\n\n"
                "-- Sample data provided by the course\n"
                "INSERT VATTU VALUES ('VT01', 'XI MANG', 'BAO', 50000, 5000);\n"
                "INSERT VATTU VALUES ('VT02', 'CAT', 'KHOI', 45000, 50000);\n"
                "INSERT VATTU VALUES ('VT03', 'GACH ONG', 'VIEN', 120, 800000);\n"
                "INSERT VATTU VALUES ('VT04', 'GACH THE', 'VIEN', 110, 800000);\n"
                "INSERT VATTU VALUES ('VT05', 'DA LON', 'KHOI', 25000, 100000);\n"
                "INSERT VATTU VALUES ('VT06', 'DA NHO', 'KHOI', 33000, 100000);\n"
                "\n"
                "INSERT KHACHHANG VALUES ('KH01', 'NGUYEN THI BE', 'TAN BINH', 8457895, 'bnt@yahoo.com');\n"
                "INSERT KHACHHANG VALUES ('KH02', 'LE HOANG NAM', 'BINH CHANH', 9878987, 'namlehoang @abc.com.vn');\n"
                "INSERT KHACHHANG VALUES ('KH03', 'TRAN THI CHIEU', 'TAN BINH', 8457895, NULL);\n"
                "INSERT KHACHHANG VALUES ('KH04', 'MAI THI QUE ANH', 'BINH CHANH ', NULL, NULL);\n"
                "INSERT KHACHHANG VALUES ('KH05', 'LE VAN SANG', 'QUAN 10', NULL, 'sanglv@hcm.vnn.vn');\n"
                "INSERT KHACHHANG VALUES ('KH06', 'TRAN HOANG KHAI', 'TAN BINH', 8457897, NULL);\n"
                "\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD001', '2000-05-12', 'KH01');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD002', '2000-05-25', 'KH02');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD003', '2000-05-25', 'KH01');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD004', '2000-05-25', 'KH04');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD005', '2000-05-26', 'KH04');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD006', '2000-05-02', 'KH03');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD007', '2000-06-22', 'KH04');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD008', '2000-06-25', 'KH03');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD009', '2000-08-15', 'KH04');\n"
                "INSERT HOADON(MaHD, Ngay, MaKH) VALUES ('HD010', '2000-08-30', 'KH01');\n"
                "\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD001', 'VT01', 5, 52000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD001', 'VT05', 10, 30000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD002', 'VT03', 10000, 150);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD003', 'VT02', 20, 55000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD004', 'VT03', 50000, 150);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD004', 'VT04', 20000, 120);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD005', 'VT05', 10, 30000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD006', 'VT04', 10000, 120);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD007', 'VT04', 20000, 125);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD008', 'VT01', 100, 55000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD009', 'VT02', 25, 48000);\n"
                "INSERT CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD010', 'VT01', 25, 57000);")

    add_heading_1(doc, "2. Programming Solutions with Student ID: QE210005")

    # Task 1
    add_heading_2(doc, "Task 1: Scalar Function QE210005_Func1")
    p_t1 = doc.add_paragraph()
    p_t1.add_run("Requirement: ").bold = True
    p_t1.add_run("Write a scalar function named QE210005_Func1 with parameter @mavt, returning the total sales revenue (SUM(SL * GiaBan)) for the corresponding material.")

    add_callout(doc, "Function Implementation (T-SQL):",
                "CREATE FUNCTION QE210005_Func1 (@mavt nvarchar(5))\n"
                "RETURNS BIGINT\n"
                "AS\n"
                "BEGIN\n"
                "    DECLARE @totalRevenue BIGINT;\n"
                "    \n"
                "    SELECT @totalRevenue = ISNULL(SUM(CAST(SL AS BIGINT) * CAST(GiaBan AS BIGINT)), 0)\n"
                "    FROM CHITIETHOADON\n"
                "    WHERE MaVT = @mavt;\n"
                "    \n"
                "    RETURN @totalRevenue;\n"
                "END;\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t1_test = doc.add_paragraph()
    p_t1_test.add_run("Execution & Test Script:\n").bold = True
    r_code1 = p_t1_test.add_run(
        "-- Test total revenue for material 'VT01' (XI MANG)\n"
        "SELECT 'VT01' AS MaVT, dbo.QE210005_Func1('VT01') AS TotalRevenue;\n"
        "-- Test for all materials:\n"
        "SELECT MaVT, TenVT, dbo.QE210005_Func1(MaVT) AS TotalRevenue FROM VATTU;"
    )
    r_code1.font.name = "Consolas"
    r_code1.font.size = Pt(9.5)
    p_t1_test.paragraph_format.space_after = Pt(8)

    # Task 2
    add_heading_2(doc, "Task 2: Scalar Function QE210005_FuncTotalHoaDon")
    p_t2 = doc.add_paragraph()
    p_t2.add_run("Requirement: ").bold = True
    p_t2.add_run("Write a scalar function to return the total amount of an invoice (SUM(SL * GiaBan)), taking @MaHD as parameter.")

    add_callout(doc, "Function Implementation (T-SQL):",
                "CREATE FUNCTION QE210005_FuncTotalHoaDon (@MaHD nvarchar(10))\n"
                "RETURNS BIGINT\n"
                "AS\n"
                "BEGIN\n"
                "    DECLARE @totalAmount BIGINT;\n"
                "    \n"
                "    SELECT @totalAmount = ISNULL(SUM(CAST(SL AS BIGINT) * CAST(GiaBan AS BIGINT)), 0)\n"
                "    FROM CHITIETHOADON\n"
                "    WHERE MaHD = @MaHD;\n"
                "    \n"
                "    RETURN @totalAmount;\n"
                "END;\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t2_test = doc.add_paragraph()
    p_t2_test.add_run("Execution & Test Script:\n").bold = True
    r_code2 = p_t2_test.add_run(
        "-- Test total amount for invoice 'HD001':\n"
        "SELECT 'HD001' AS MaHD, dbo.QE210005_FuncTotalHoaDon('HD001') AS InvoiceTotal;\n"
        "-- Check all invoices:\n"
        "SELECT MaHD, Ngay, MaKH, dbo.QE210005_FuncTotalHoaDon(MaHD) AS CalculatedTotal FROM HOADON;"
    )
    r_code2.font.name = "Consolas"
    r_code2.font.size = Pt(9.5)
    p_t2_test.paragraph_format.space_after = Pt(8)

    # Task 3
    add_heading_2(doc, "Task 3: Stored Procedure QE210005_Proc1 (Update Customer Address)")
    p_t3 = doc.add_paragraph()
    p_t3.add_run("Requirement: ").bold = True
    p_t3.add_run("Write a stored procedure named QE210005_Proc1 with parameters @makh and @diachi. This procedure updates the address of the corresponding customer, verifying customer existence.")

    add_callout(doc, "Stored Procedure Implementation (T-SQL):",
                "CREATE PROCEDURE QE210005_Proc1\n"
                "    @makh nvarchar(5),\n"
                "    @diachi nvarchar(50)\n"
                "AS\n"
                "BEGIN\n"
                "    SET NOCOUNT ON;\n"
                "    \n"
                "    IF NOT EXISTS (SELECT 1 FROM KHACHHANG WHERE MaKH = @makh)\n"
                "    BEGIN\n"
                "        RAISERROR(N'Error: Customer code %s does not exist in KHACHHANG table!', 16, 1, @makh);\n"
                "        RETURN -1;\n"
                "    END;\n"
                "    \n"
                "    UPDATE KHACHHANG\n"
                "    SET DiaChi = @diachi\n"
                "    WHERE MaKH = @makh;\n"
                "    \n"
                "    PRINT N'Success: Address updated successfully for customer: ' + @makh;\n"
                "    RETURN 0;\n"
                "END;\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t3_test = doc.add_paragraph()
    p_t3_test.add_run("Execution & Test Script:\n").bold = True
    r_code3 = p_t3_test.add_run(
        "-- Update address for customer KH01:\n"
        "EXEC QE210005_Proc1 @makh = 'KH01', @diachi = N'123 Le Loi, Quan 1, TP.HCM';\n"
        "-- Verify update:\n"
        "SELECT * FROM KHACHHANG WHERE MaKH = 'KH01';"
    )
    r_code3.font.name = "Consolas"
    r_code3.font.size = Pt(9.5)
    p_t3_test.paragraph_format.space_after = Pt(8)

    # Task 4
    add_heading_2(doc, "Task 4: Stored Procedure QE210005_ProcAddItem (Add Item to Invoice)")
    p_t4 = doc.add_paragraph()
    p_t4.add_run("Requirement: ").bold = True
    p_t4.add_run("Write a stored procedure to add an item into an invoice (CHITIETHOADON). It ensures invoice and material validity, verifies stock sufficiency, updates or inserts line items, and deducts inventory quantity.")

    add_callout(doc, "Stored Procedure Implementation (T-SQL):",
                "CREATE PROCEDURE QE210005_ProcAddItem\n"
                "    @MaHD nvarchar(10),\n"
                "    @MaVT nvarchar(5),\n"
                "    @SL int,\n"
                "    @KhuyenMai int = 0,\n"
                "    @GiaBan int = NULL\n"
                "AS\n"
                "BEGIN\n"
                "    SET NOCOUNT ON;\n"
                "    DECLARE @CurrentStock int, @DefaultPrice int;\n"
                "\n"
                "    -- 1. Validate input BEFORE the transaction (outside TRY, so RETURN codes are reachable)\n"
                "    IF @SL IS NULL OR @SL <= 0\n"
                "    BEGIN\n"
                "        RAISERROR(N'Error: Quantity must be greater than 0!', 16, 1);\n"
                "        RETURN -4;\n"
                "    END;\n"
                "\n"
                "    IF NOT EXISTS (SELECT 1 FROM HOADON WHERE MaHD = @MaHD)\n"
                "    BEGIN\n"
                "        RAISERROR(N'Error: Invoice %s does not exist!', 16, 1, @MaHD);\n"
                "        RETURN -1;\n"
                "    END;\n"
                "\n"
                "    SELECT @CurrentStock = SLTon, @DefaultPrice = GiaMua\n"
                "    FROM VATTU WHERE MaVT = @MaVT;\n"
                "\n"
                "    IF @CurrentStock IS NULL\n"
                "    BEGIN\n"
                "        RAISERROR(N'Error: Material %s does not exist!', 16, 1, @MaVT);\n"
                "        RETURN -2;\n"
                "    END;\n"
                "\n"
                "    IF @CurrentStock < @SL\n"
                "    BEGIN\n"
                "        RAISERROR(N'Error: Insufficient stock for material %s! Available: %d, Requested: %d', 16, 1, @MaVT, @CurrentStock, @SL);\n"
                "        RETURN -3;\n"
                "    END;\n"
                "\n"
                "    IF @GiaBan IS NULL\n"
                "        SET @GiaBan = CAST(@DefaultPrice * 1.15 AS INT); -- Default 15% markup\n"
                "\n"
                "    -- 2. Insert/update the line item and deduct stock as ONE atomic transaction\n"
                "    BEGIN TRY\n"
                "        BEGIN TRANSACTION;\n"
                "\n"
                "        IF EXISTS (SELECT 1 FROM CHITIETHOADON WHERE MaHD = @MaHD AND MaVT = @MaVT)\n"
                "            UPDATE CHITIETHOADON\n"
                "            SET SL = SL + @SL, GiaBan = @GiaBan, KhuyenMai = @KhuyenMai\n"
                "            WHERE MaHD = @MaHD AND MaVT = @MaVT;\n"
                "        ELSE\n"
                "            INSERT INTO CHITIETHOADON(MaHD, MaVT, SL, KhuyenMai, GiaBan)\n"
                "            VALUES (@MaHD, @MaVT, @SL, @KhuyenMai, @GiaBan);\n"
                "\n"
                "        UPDATE VATTU\n"
                "        SET SLTon = SLTon - @SL\n"
                "        WHERE MaVT = @MaVT;\n"
                "\n"
                "        COMMIT TRANSACTION;\n"
                "    END TRY\n"
                "    BEGIN CATCH\n"
                "        IF @@TRANCOUNT > 0 ROLLBACK TRANSACTION;\n"
                "        THROW; -- re-raise the original error\n"
                "    END CATCH;\n"
                "\n"
                "    PRINT N'Success: Added material ' + @MaVT + N' to invoice ' + @MaHD;\n"
                "    RETURN 0;\n"
                "END;\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t4_test = doc.add_paragraph()
    p_t4_test.add_run("Execution & Test Script:\n").bold = True
    r_code4 = p_t4_test.add_run(
        "-- Add 10 units of VT01 to HD002 at price 55000:\n"
        "EXEC QE210005_ProcAddItem @MaHD = 'HD002', @MaVT = 'VT01', @SL = 10, @GiaBan = 55000;\n"
        "-- Verify invoice details and inventory reduction:\n"
        "SELECT * FROM CHITIETHOADON WHERE MaHD = 'HD002';\n"
        "SELECT MaVT, TenVT, SLTon FROM VATTU WHERE MaVT = 'VT01';"
    )
    r_code4.font.name = "Consolas"
    r_code4.font.size = Pt(9.5)
    p_t4_test.paragraph_format.space_after = Pt(8)

    # Task 5
    add_heading_2(doc, "Task 5: Trigger QE210005_Trig1 on Table CHITIETHOADON")
    p_t5 = doc.add_paragraph()
    p_t5.add_run("Requirement: ").bold = True
    p_t5.add_run("Add a Tongtien column into table HOADON. Write a trigger named QE210005_Trig1 on table CHITIETHOADON so that whenever data is inserted, updated, or deleted in CHITIETHOADON, the trigger automatically recalculates and updates Tongtien in HOADON (Tongtien = SUM(SL * GiaBan)).")

    add_callout(doc, "Schema Alteration & Trigger Implementation (T-SQL):",
                "-- Step 1: Add Tongtien column to HOADON table\n"
                "IF NOT EXISTS (\n"
                "    SELECT 1 FROM sys.columns \n"
                "    WHERE object_id = OBJECT_ID('HOADON') AND name = 'Tongtien'\n"
                ")\n"
                "BEGIN\n"
                "    ALTER TABLE HOADON ADD Tongtien BIGINT NULL;\n"
                "END;\n"
                "GO\n\n"
                "-- Step 2: Implement robust multi-row safe trigger\n"
                "CREATE TRIGGER QE210005_Trig1\n"
                "ON CHITIETHOADON\n"
                "AFTER INSERT, UPDATE, DELETE\n"
                "AS\n"
                "BEGIN\n"
                "    SET NOCOUNT ON;\n"
                "    \n"
                "    -- Collect all distinct invoice IDs affected by insert, update, or delete\n"
                "    DECLARE @AffectedInvoices TABLE (MaHD nvarchar(10));\n"
                "    \n"
                "    INSERT INTO @AffectedInvoices (MaHD)\n"
                "    SELECT DISTINCT MaHD FROM inserted\n"
                "    UNION\n"
                "    SELECT DISTINCT MaHD FROM deleted;\n"
                "    \n"
                "    -- Recalculate and update Tongtien for all affected invoices\n"
                "    UPDATE H\n"
                "    SET H.Tongtien = (\n"
                "        SELECT ISNULL(SUM(CAST(C.SL AS BIGINT) * CAST(C.GiaBan AS BIGINT)), 0)\n"
                "        FROM CHITIETHOADON C\n"
                "        WHERE C.MaHD = H.MaHD\n"
                "    )\n"
                "    FROM HOADON H\n"
                "    INNER JOIN @AffectedInvoices A ON H.MaHD = A.MaHD;\n"
                "END;\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t5_test = doc.add_paragraph()
    p_t5_test.add_run("Execution & Test Script:\n").bold = True
    r_code5 = p_t5_test.add_run(
        "-- Initialize Tongtien for existing data:\n"
        "UPDATE H SET H.Tongtien = (SELECT ISNULL(SUM(CAST(SL AS BIGINT)*CAST(GiaBan AS BIGINT)), 0) FROM CHITIETHOADON C WHERE C.MaHD = H.MaHD) FROM HOADON H;\n"
        "-- Test Trigger on insert:\n"
        "INSERT INTO CHITIETHOADON(MaHD, MaVT, SL, GiaBan) VALUES ('HD001', 'VT02', 2, 60000);\n"
        "-- Inspect automatic update:\n"
        "SELECT MaHD, Ngay, MaKH, Tongtien FROM HOADON WHERE MaHD = 'HD001';"
    )
    r_code5.font.name = "Consolas"
    r_code5.font.size = Pt(9.5)
    p_t5_test.paragraph_format.space_after = Pt(8)

    # Task 6
    add_heading_2(doc, "Task 6: View QE210005_View1 (Customers Buying GACH ONG)")
    p_t6 = doc.add_paragraph()
    p_t6.add_run("Requirement: ").bold = True
    p_t6.add_run("Write a view named QE210005_View1 to extract the list of distinct customers who purchased material 'GACH ONG'.")

    add_callout(doc, "View Implementation (T-SQL):",
                "CREATE VIEW QE210005_View1\n"
                "AS\n"
                "SELECT DISTINCT \n"
                "    K.MaKH,\n"
                "    K.TenKH,\n"
                "    K.DiaChi,\n"
                "    K.DT,\n"
                "    K.Email\n"
                "FROM KHACHHANG K\n"
                "INNER JOIN HOADON H ON K.MaKH = H.MaKH\n"
                "INNER JOIN CHITIETHOADON C ON H.MaHD = C.MaHD\n"
                "INNER JOIN VATTU V ON C.MaVT = V.MaVT\n"
                "WHERE V.TenVT = N'GACH ONG';\n"
                "GO",
                bg_hex=HEX_LIGHT_BLUE, border_color="003366")

    p_t6_test = doc.add_paragraph()
    p_t6_test.add_run("Execution & Test Script:\n").bold = True
    r_code6 = p_t6_test.add_run(
        "-- Query the view:\n"
        "SELECT * FROM QE210005_View1;\n"
        "-- Check result matches customers KH02 and KH04 who bought VT03 (GACH ONG)."
    )
    r_code6.font.name = "Consolas"
    r_code6.font.size = Pt(9.5)
    p_t6_test.paragraph_format.space_after = Pt(10)

    # Summary
    add_heading_1(doc, "3. Summary of Programmed Database Objects")
    tbl_prog_sum = doc.add_table(rows=7, cols=4)
    style_standard_table(
        tbl_prog_sum,
        [Inches(1.0), Inches(2.2), Inches(1.5), Inches(2.0)],
        ["Task #", "Object Name", "Type", "Operational Purpose"],
        [
            ["Task 1", "QE210005_Func1", "Scalar Function", "Total sales revenue for material @mavt"],
            ["Task 2", "QE210005_FuncTotalHoaDon", "Scalar Function", "Total calculated amount of invoice @MaHD"],
            ["Task 3", "QE210005_Proc1", "Stored Procedure", "Update customer address with validation"],
            ["Task 4", "QE210005_ProcAddItem", "Stored Procedure", "Add line item & deduct inventory safely"],
            ["Task 5", "QE210005_Trig1", "DML Trigger (AFTER)", "Auto-calculate Tongtien on insert/update/delete"],
            ["Task 6", "QE210005_View1", "Relational View", "Extract distinct customers buying 'GACH ONG'"]
        ]
    )

    doc.save(output_path)
    print(f"Lab 5 successfully generated at: {output_path}")

if __name__ == "__main__":
    os.makedirs("output", exist_ok=True)
    generate_lab4("output/Lab4.docx")
    generate_lab5("output/Lab5.docx")
    print("LAB 4 & LAB 5 SUCCESSFULLY CREATED IN output/ DIRECTORY!")
